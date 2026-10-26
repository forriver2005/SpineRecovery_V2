from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile, ZIP_DEFLATED
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re,hashlib,json,math

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parent.parent/'output'/'patent_direction3'
OUT.mkdir(parents=True,exist_ok=True)
REF=Path('C:/Users/23548/Documents/xwechat_files/wxid_dv2i6u8f201q12_fa7a/msg/file/2026-10/一种基于动作偏差定位与限时振动反馈的肌骨康复训练引导方法及系统.docx')
FONT='C:/Windows/Fonts/simsun.ttc'

class Diagram:
 def __init__(self,w,h):
  self.im=Image.new('RGB',(w,h),'white');self.d=ImageDraw.Draw(self.im)
 def text(self,x,y,s,size=44):
  f=ImageFont.truetype(FONT,size);lines=s.split('\n');height=len(lines)*(size+12)-12
  for i,line in enumerate(lines):
   bb=self.d.textbbox((0,0),line,font=f);self.d.text((x-(bb[2]-bb[0])/2,y-height/2+i*(size+12)-bb[1]),line,font=f,fill='black')
 def box(self,xy,s,size=44):
  self.d.rectangle(xy,outline='black',width=3);self.text((xy[0]+xy[2])/2,(xy[1]+xy[3])/2,s,size)
 def arrow(self,pts,label=None,lp=None):
  self.d.line(pts,fill='black',width=3);a,b=pts[-2:];ang=math.atan2(b[1]-a[1],b[0]-a[0]);l=18
  p1=(b[0]-l*math.cos(ang-.45),b[1]-l*math.sin(ang-.45));p2=(b[0]-l*math.cos(ang+.45),b[1]-l*math.sin(ang+.45));self.d.polygon([b,p1,p2],fill='black')
  if label:self.text(*lp,label,36)
 def save(self,n):self.im.save(ROOT/f'fig{n}.png',dpi=(300,300))

g=Diagram(1200,1460)
labels=['S1  获取动作描述符与处方','S2  分类并构建训练原语图','S3  检查模板兼容性\n写入处方并绑定任务','S4  处理运动数据并驱动游戏\n检查约束与节点条件','S5  提交有效完成事件\n更新剂量并调度或结束']
for i,t in enumerate(labels):
 y=50+i*280;g.box((80,y,1120,y+190),t,48)
 if i<4:g.arrow([(600,y+190),(600,y+280)])
g.save(1)

g=Diagram(1500,1510)
g.box((30,30,710,260),'动作描述符\n体位  通道  阶段  协同\n稳定约束与传感器需求',42)
g.box((790,30,1470,260),'处方参数\n目标范围  次数  组数\n保持  休息  侧别计数',42)
g.arrow([(370,260),(370,325),(750,325),(750,390)])
g.arrow([(1130,260),(1130,325),(750,325)])
g.box((170,390,1330,650),'训练原语图\n准备 → 到达 → 保持 → 返回\n并行完成与交替关系  重复与休息',46)
g.arrow([(750,650),(750,750)])
g.box((170,750,1330,980),'兼容性检查与绑定\n通道类型  必需节点  协同条件\n测量能力  计时方式  暂停与休息',46)
g.arrow([(750,980),(750,1020),(370,1020),(370,1080)])
g.arrow([(750,1020),(1130,1020),(1130,1080)])
g.box((30,1080,710,1370),'游戏模板实例\n连续控制量驱动画面\n有效事件推进任务',44)
g.box((790,1080,1470,1370),'独立剂量记录\n组与侧别分别记录\n事件去重及剩余剂量',44)
g.save(2)

g=Diagram(1500,1630)
boxes=[('准备 READY\n确认起始区间后离开',30),('到达 REACH\n各必需通道与协同条件满足',280),('保持 HOLD\n连续有效时间达到处方时长',530),('返回 RETURN\n全部必需通道返回并确认',780),('提交本次有效完成\n去重后更新组与侧别剂量',1030),('本组完成：休息 REST\n或直接结束 END',1320)]
for i,(t,y) in enumerate(boxes):
 g.box((110,y,1060,y+170),t,43)
 if i<5:g.arrow([(585,y+170),(585,boxes[i+1][1])])
g.box((1140,480,1475,760),'保持中断\n本次连续\n计时清零',40)
g.arrow([(1060,600),(1140,600)]);g.arrow([(1300,480),(1300,460),(1090,460),(1090,490),(1000,490),(1000,530)])
g.arrow([(110,1115),(45,1115),(45,115),(110,115)])
g.text(290,1245,'组内继续或切换组合',35)
g.box((110,1540,1475,1620),'数据失效或停止：暂停，撤销未提交周期；恢复后重新准备',35)
g.save(3)

g=Diagram(1500,1430)
g.box((30,30,710,280),'鸟狗式\n四点支撑\n活动肢体与支撑约束',46)
g.box((790,30,1470,280),'死虫式\n仰卧\n独立角度定义与稳定约束',46)
g.arrow([(370,280),(370,360),(750,360),(750,430)])
g.arrow([(1130,280),(1130,360),(750,360)])
g.box((90,430,1410,690),'共同控制结构\n右上肢与左下肢  ／  左上肢与右下肢\n组合内并行完成  组合间返回后交替',46)
g.arrow([(750,690),(750,790)])
g.box((90,790,1410,1000),'同一对侧交替类游戏模板\n独立绑定各自动作特征与处方参数\n两通道均满足后保持  返回后提交',46)
g.arrow([(750,1000),(750,1080),(370,1080),(370,1150)])
g.arrow([(750,1080),(1130,1080),(1130,1150)])
g.box((30,1150,710,1380),'示例处方一\n每组合6次 × 2组\n每次保持5 s',44)
g.box((790,1150,1470,1380),'示例处方二\n每组合4次 × 1组\n每次保持3 s',44)
g.save(4)

g=Diagram(1500,1450)
g.box((90,30,1040,200),'登记新动作描述符\n特征规则  体位  阶段  约束',44)
g.box((90,300,1040,470),'字段校验与控制结构分类\n构建训练原语图',44)
g.box((90,570,1040,780),'检查测量能力与模板接口\n全部必需节点及约束均可执行',44)
g.box((90,900,1040,1070),'通过后绑定已有模板\n按处方生成任务实例',44)
g.box((90,1190,1040,1360),'运行与剂量提交\n记录版本及有效完成事件',44)
for a,b in [(200,300),(470,570),(780,900),(1070,1190)]:g.arrow([(565,a),(565,b)])
g.box((1130,560,1470,870),'不通过\n输出缺失项\n停止启动',40)
g.arrow([(1040,675),(1130,675)])
g.text(1380,965,'补充后',36);g.text(1380,1015,'重新检查',36)
g.arrow([(1300,870),(1300,1100),(1080,1100),(1080,420),(1040,420)])
g.save(5)

text=(ROOT/'draft.txt').read_text(encoding='utf8')
title=text.split('@TITLE\n')[1].split('@ABSTRACT')[0].strip()
abstract=text.split('@ABSTRACT\n')[1].split('@CLAIMS')[0].strip()
claims=text.split('@CLAIMS\n')[1].split('@DESCRIPTION')[0].strip().splitlines()
description=text.split('@DESCRIPTION\n')[1].strip().splitlines()
doc=Document(REF)
sects=[deepcopy(s._sectPr) for s in doc.sections]
body=doc._element.body
for e in list(body):body.remove(e)
normal=doc.styles['Normal']
normal.font.color.rgb=RGBColor(0,0,0)
heading=doc.styles['Heading 1']
for b in heading._element.xpath('./w:pPr/w:pBdr'):b.getparent().remove(b)
heading.font.color.rgb=RGBColor(0,0,0)
try: ts=doc.styles['Title']
except KeyError:ts=doc.styles.add_style('Title',1)
ts.base_style=normal;ts.font.size=Pt(14);ts.font.bold=True;ts.font.color.rgb=RGBColor(0,0,0)
ts.font.name='Times New Roman';ts._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
ts.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.CENTER
ts.paragraph_format.space_after=Pt(12);ts.paragraph_format.keep_with_next=True

def para(t,bold=False,center=False,small=False,keep=False):
 p=doc.add_paragraph(style='Normal');p.paragraph_format.line_spacing=1.5
 p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
 p.paragraph_format.first_line_indent=Pt(24) if not center and not bold and not small else Pt(0)
 p.paragraph_format.widow_control=True;p.paragraph_format.keep_with_next=keep
 p.alignment=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
 if small:p.paragraph_format.line_spacing=1.2;p.paragraph_format.space_after=Pt(4)
 # Turn DOI text into a clickable hyperlink while preserving printed citation format.
 m=re.search(r'DOI: ([^ ]+)\.$',t)
 normal_text=t[:m.start()] if m else t
 r=p.add_run(normal_text);r.bold=bold;r.font.size=Pt(10.5 if small else 12)
 if m:
  p.add_run('DOI: ').font.size=Pt(10.5)
  url='https://doi.org/'+m.group(1)
  rid=p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True)
  h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rid);r=OxmlElement('w:r');rp=OxmlElement('w:rPr');sz=OxmlElement('w:sz');sz.set(qn('w:val'),'21');rp.append(sz);r.append(rp);tx=OxmlElement('w:t');tx.text=m.group(1);r.append(tx);h.append(r);p._p.append(h);p.add_run('.').font.size=Pt(10.5)
 return p
def chapter(t):
 p=doc.add_paragraph(t,style='Heading 1');p.paragraph_format.space_after=Pt(12);p.paragraph_format.space_before=Pt(0)
 return p
def section_end(idx):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1;p._p.get_or_add_pPr().append(deepcopy(sects[idx]))
def fig(n,width=15):
 p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(4)
 p.add_run().add_picture(str(ROOT/f'fig{n}.png'),width=Cm(width));para('图'+str(n),center=True)

chapter('说   明   书   摘   要');para(abstract);section_end(0)
chapter('摘   要   附   图');para('图1作为摘要附图。',center=True);fig(1,10);section_end(1)
chapter('权   利   要   求   书')
for line in claims:para(line)
section_end(2)
chapter('说   明   书');doc.add_paragraph(title,style='Title')
for line in description:
 if line.startswith('##'):
  p=para(line[2:],bold=True,keep=True);p.paragraph_format.space_before=Pt(8)
 elif line.startswith('#'):
  p=para(line[1:],bold=True,keep=True);p.paragraph_format.space_before=Pt(8)
 elif line.startswith('='):para(line[1:],center=True)
 elif re.match(r'^\[\d+\]',line):para(line,small=True)
 else:para(line)
doc.add_page_break();chapter('说   明   书   附   图')
for n in range(1,6):
 if n>1:doc.add_page_break()
 fig(n,14.5)
body.append(deepcopy(sects[3]))
doc.core_properties.title=title;doc.core_properties.subject='康复处方与游戏动作映射';doc.core_properties.author='';doc.core_properties.last_modified_by=''
inter=ROOT/'assembled.docx';doc.save(inter)
final=OUT/'处方游戏动作映射算法_专利申请文件_含文献引用_v1.docx'
editable={'word/document.xml','word/_rels/document.xml.rels','word/styles.xml','word/settings.xml','docProps/core.xml','docProps/app.xml','[Content_Types].xml'}
with ZipFile(REF) as zin,ZipFile(inter) as gen,ZipFile(final,'w',ZIP_DEFLATED) as zout:
 orig=set(zin.namelist())
 for name in zin.namelist():zout.writestr(name,gen.read(name) if name in editable else zin.read(name))
 for name in gen.namelist():
  if name not in orig:zout.writestr(name,gen.read(name))
with ZipFile(REF) as a,ZipFile(final) as b:
 mismatches=[n for n in a.namelist() if n not in editable and a.read(n)!=b.read(n)]
 assert not mismatches,mismatches
assert len(abstract)<=300
assert len(Document(final).sections)==4
print(json.dumps({'output':str(final),'abstract_chars':len(abstract),'claims':len([x for x in claims if re.match(r'^\d+、',x)]),'figures':5,'references':6,'characters':len(text),'preserved_parts':len(orig-editable)},ensure_ascii=False))
