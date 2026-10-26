from pathlib import Path
from zipfile import ZipFile
from lxml import etree
from docx import Document
import json,hashlib
b=Path('C:/Users/23548/Documents/xwechat_files/wxid_dv2i6u8f201q12_fa7a/msg/file/2026-10');out=Path('tmp/patent_direction3')
files=list(b.glob('增强现实镜像康复训练指导系统及方法_专利申请文件_v2(2).docx'))+list(b.glob('一种基于动作偏差定位与限时振动反馈的肌骨康复训练引导方法及系统.docx'))
for i,p in enumerate(files,1):
 d=Document(p)
 print('REF',i,'sections',len(d.sections),'paragraphs',len(d.paragraphs),'tables',len(d.tables),'images',len(d.inline_shapes))
 for s in d.sections: print('GEOMETRY',*[round(getattr(s,k).cm,3) for k in ['page_width','page_height','top_margin','bottom_margin','left_margin','right_margin']])
 for j in [0,1,8,10,20,35]:
  pa=d.paragraphs[j];print('P',j,pa.text[:65], 'STYLE',pa.style.name,'ALIGN',pa.alignment,'SPACE',pa.paragraph_format.line_spacing,'INDENT',pa.paragraph_format.first_line_indent)
  print([(r.font.name,r.font.size.pt if r.font.size else None,r.bold) for r in pa.runs[:3]])
 with ZipFile(p) as z:
  root=etree.fromstring(z.read('word/document.xml'));ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
  print('EQUATIONS',root.xpath('//m:oMath//m:t/text()',namespaces=ns))
  inventory=[{'path':n,'size':len(z.read(n)),'sha256':hashlib.sha256(z.read(n)).hexdigest()} for n in z.namelist()]
  (out/f'ref{i}-inventory.json').write_text(json.dumps(inventory),encoding='utf8')
